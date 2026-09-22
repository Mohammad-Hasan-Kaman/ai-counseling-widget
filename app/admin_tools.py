# -*- coding: utf-8 -*-
"""
ماژول ابزارهای پنل ادمین - مرکز مشاوره خانواده نیک‌روان
شامل: خروجی اکسل دیتای کاربران و ارسال پیام برادکست (همگانی)
این ماژول مستقل است و هیچ تغییری در کدهای موجود ایجاد نمی‌کند.
"""
import io
import json
import asyncio
import logging

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.config import USER_RECORDS_DB

log = logging.getLogger(__name__)

BROADCAST_BATCH_DELAY = 0.05   # فاصله بین ارسال‌ها برای جلوگیری از محدودیت سرور بله
BROADCAST_REPORT_EVERY = 20    # هر چند نفر یک‌بار گزارش پیشرفت ارسال شود


# ══════════════════════════════════════════════
#  بخش ۱: خروجی اکسل دیتای کاربران
# ══════════════════════════════════════════════

_HEADER_FILL = PatternFill("solid", fgColor="2E5E4E")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

USERS_HEADERS = REQUESTS_HEADERS = [
    "ردیف", "نام و نام خانوادگی", "شماره تماس", "شناسه بله (Chat ID)",
    "جنسیت", "سن", "موضوع مشاوره (مشکل اصلی)",
    "سابقه مراجعه قبلی به روانشناس/روانپزشک", "توضیحات درمان قبلی",
    "انتظار یا رویکرد موردنظر از مشاوره", "ترجیح جنسیت مشاور",
    "شعبه انتخابی",
    "GHQ - جسمانی‌شدن (سوماتیک)", "GHQ - اضطراب و بی‌خوابی",
    "GHQ - اختلال کارکرد اجتماعی", "GHQ - افسردگی",
    "نتیجه نهایی GHQ (امتیاز کل)", "وضعیت GHQ",
    "مشاوران پیشنهادی شده", "دلیل پیشنهاد", "شماره مراجعه", "تاریخ ثبت درخواست",
    "پاسخ سؤال‌های سفارشی",
]


def _style_header(ws, headers):
    fill, font, align = _HEADER_FILL, _HEADER_FONT, _HEADER_ALIGN
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.freeze_panes = "A2"
    for col_idx in range(1, len(headers) + 1):
        width = max(14, min(30, len(headers[col_idx - 1]) * 1.6 + 6))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _fmt_recs(recommendations_json: str | None) -> tuple[str, str]:
    """استخراج نام مشاوران و دلایل پیشنهاد از JSON ذخیره‌شده"""
    if not recommendations_json:
        return "", ""
    try:
        import json
        recs = json.loads(recommendations_json)
        if isinstance(recs, list):
            names = " | ".join(r.get("name", "") for r in recs if isinstance(r, dict))
            reasons = " | ".join(
                f"{r.get('name','')}: {r.get('reason','')}" for r in recs if isinstance(r, dict)
            )
            return names, reasons
        return str(recs), ""
    except Exception:
        return "", ""


def _ghq_status(total) -> str:
    """برچسب وضعیت بر اساس امتیاز کل GHQ (آستانه شدت: ۴۳)"""
    if total is None:
        return "انجام نداده"
    if total >= 43:
        return "شدید"
    if total >= 24:
        return "متوسط"
    return "طبیعی"


def export_users_excel(tenant_id: int = 1) -> bytes:
    """
    تولید خروجی اکسل دیتای کاربران در حافظه (بدون ساخت فایل روی دیسک).
    هر دو شیت کامل و یکسان‌اند: شیت ۱ «کاربران» (مرتب بر اساس نام) و
    شیت ۲ «سوابق درخواست‌ها» (مرتب بر اساس تاریخ، جدیدترین اول) —
    هر ردیف = یک درخواست با تمام اطلاعات پرسیده‌شده از کاربر.
    خروجی: بایت‌های فایل xlsx
    """
    import sqlite3
    conn = sqlite3.connect(str(USER_RECORDS_DB))

    requests_rows_name = conn.execute("""
        SELECT full_name, phone, telegram_id, gender, age, topic,
               has_prev_therapy, prev_detail, expectation, preferred_gender,
               branch, ghq_scores, ghq_total, recommendations, request_number, created_at, custom_data
        FROM user_requests WHERE tenant_id=? ORDER BY full_name COLLATE NOCASE ASC, request_number ASC
    """, (tenant_id,)).fetchall()

    requests_rows_date = conn.execute("""
        SELECT full_name, phone, telegram_id, gender, age, topic,
               has_prev_therapy, prev_detail, expectation, preferred_gender,
               branch, ghq_scores, ghq_total, recommendations, request_number, created_at, custom_data
        FROM user_requests WHERE tenant_id=? ORDER BY created_at DESC
    """, (tenant_id,)).fetchall()
    conn.close()

    wb = Workbook()

    def _safe_cell(v):
        """جلوگیری از formula injection در اکسل: مقادیر متنی که با = + - @ شروع شوند"""
        if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
            return "'" + v
        return v

    def _fmt_custom(custom_json: str | None) -> str:
        """پاسخ سؤال‌های سفارشی فلو (q_*) به‌صورت «کلید: مقدار | ...»"""
        if not custom_json:
            return ""
        try:
            d = json.loads(custom_json)
            if not isinstance(d, dict) or not d:
                return ""
            return " | ".join(f"{k}: {v}" for k, v in d.items())
        except Exception:
            return ""

    def _write_request_row(ws, i, r):
        (name, phone, tg_id, gender, age, topic, has_prev, prev_detail,
         expectation, pref_gender, branch, ghq_json, ghq_total, recs,
         req_number, created, custom_json) = r

        # استخراج زیرمقیاس‌های GHQ از JSON ذخیره‌شده
        somatic = anxiety = social = depression = None
        if ghq_json:
            try:
                g = json.loads(ghq_json)
                somatic = g.get("somatic")
                anxiety = g.get("anxiety")
                social = g.get("social")
                depression = g.get("depression")
                if ghq_total is None and isinstance(g.get("total"), int):
                    ghq_total = g["total"]
            except Exception:
                pass

        rec_names, rec_reasons = _fmt_recs(recs)

        def _clean(v):
            """تبدیل None و رشته 'None' به رشته خالی برای نمایش تمیز در اکسل"""
            if v is None:
                return ""
            s = str(v)
            return "" if s.strip().lower() == "none" else s

        ws.append([_safe_cell(v) for v in [
            i, name, phone, tg_id or "",
            gender or "",
            age if age is not None else "",
            topic or "",
            "بله" if has_prev else "خیر",
            _clean(prev_detail),
            expectation or "",
            pref_gender or "",
            branch or "-",
            somatic if somatic is not None else "-",
            anxiety if anxiety is not None else "-",
            social if social is not None else "-",
            depression if depression is not None else "-",
            ghq_total if ghq_total is not None else "-",
            _ghq_status(ghq_total),
            rec_names,
            rec_reasons,
            req_number if req_number is not None else "",
            created or "",
            _fmt_custom(custom_json),
        ]])

    # ── شیت ۱: کاربران (همه اطلاعات، مرتب بر اساس نام) ──
    ws_users = wb.active
    ws_users.title = "کاربران"
    _style_header(ws_users, USERS_HEADERS)
    for i, r in enumerate(requests_rows_name, start=1):
        _write_request_row(ws_users, i, r)

    # ── شیت ۲: سوابق درخواست‌ها (همان ستون‌ها، مرتب بر اساس تاریخ) ──
    ws_req = wb.create_sheet("سوابق درخواست‌ها")
    _style_header(ws_req, REQUESTS_HEADERS)
    for i, r in enumerate(requests_rows_date, start=1):
        _write_request_row(ws_req, i, r)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
